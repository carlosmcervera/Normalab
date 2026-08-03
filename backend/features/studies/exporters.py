"""Exportadores del estudio: informe en DOCX y tablas de cálculo en XLSX.

- build_docx(report, study, project): reproduce el informe completo (mismas
  secciones que el PDF en pantalla) en un documento Word editable.
- build_xlsx(study, project): libro con hojas Superficies, Ocupación y una hoja
  por sector con la carga de fuego (Qs) escrita como FÓRMULAS vivas, de modo que
  el técnico pueda editar los datos y recalcular en Excel.
"""
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# --------------------------------------------------------------------------- #
# Helpers de formato
# --------------------------------------------------------------------------- #
def _fmt(v, dec=0):
    if v is None:
        return "—"
    if isinstance(v, (int, float)):
        s = f"{v:,.{dec}f}"
        return s.replace(",", "§").replace(".", ",").replace("§", ".")
    return str(v)


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #
ORANGE = RGBColor(0xEA, 0x58, 0x0C)
GREY = RGBColor(0x64, 0x74, 0x8B)

ESC_TIPO_LABEL = {
    "abierta": "abierta", "protegida": "protegida",
    "especialmente_protegida": "especialmente protegida", "compartimentada": "compartimentada",
}


def _shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def _set_cell(cell, text, bold=False, align=None, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(8.5)
    if color:
        run.font.color.rgb = color


def _ref_table(doc, t):
    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run(t.get("titulo", ""))
    r.bold = True
    r.font.size = Pt(9.5)
    legend = t.get("legend") or []
    color_by_idx = {l["idx"]: l["color"].lstrip("#").upper() for l in legend}
    if legend:
        lg = doc.add_paragraph()
        for i, l in enumerate(legend):
            run = lg.add_run(("  " if i else "") + f"■ {l['nombre']}")
            run.font.size = Pt(8)
            try:
                run.font.color.rgb = RGBColor.from_string(l["color"].lstrip("#").upper())
            except Exception:
                pass
    headers = t.get("headers", [])
    rows = t.get("rows", [])
    cell_sectors = t.get("cell_sectors") or {}
    row_sectors = t.get("row_sectors") or {}
    hl_rows = set(t.get("highlight_rows") or [])
    hl_cells = {tuple(c) for c in (t.get("highlight_cells") or [])}
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        _set_cell(table.rows[0].cells[i], h, bold=True)
        _shade(table.rows[0].cells[i], "E2E8F0")
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        rsec = row_sectors.get(str(ri))
        for ci, val in enumerate(row):
            idxs = cell_sectors.get(f"{ri}-{ci}")
            if not idxs and rsec and ci == 0:
                idxs = rsec
            shade = None
            bold = False
            if idxs:
                shade = color_by_idx.get(idxs[0], "FDBA74")
                bold = True
            elif rsec:
                shade = "FEF3C7"
            elif ri in hl_rows or (ri, ci) in hl_cells:
                shade = "FDBA74"
                bold = True
            _set_cell(cells[ci], val, bold=bold)
            if shade:
                _shade(cells[ci], shade)
    if t.get("nota"):
        n = doc.add_paragraph()
        rn = n.add_run(t["nota"])
        rn.italic = True
        rn.font.size = Pt(8)
    ref = doc.add_paragraph()
    rr = ref.add_run(f"Fuente: {t.get('referencia', '')}")
    rr.italic = True
    rr.font.size = Pt(8)
    rr.font.color.rgb = GREY


def _kv_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for row in rows:
        cells = table.add_row().cells
        _set_cell(cells[0], row.get("k", ""), color=GREY)
        _set_cell(cells[1], row.get("v", ""), bold=True)


def _generic_table(doc, headers, rows, shade_header=True, totals=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        _set_cell(table.rows[0].cells[i], h, bold=True)
        if shade_header:
            _shade(table.rows[0].cells[i], "E2E8F0")
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            _set_cell(cells[i], val)
    if totals:
        cells = table.add_row().cells
        for i, val in enumerate(totals):
            _set_cell(cells[i], val, bold=True)
            _shade(cells[i], "F1F5F9")


def _sector_docx(doc, s):
    nr = s.get("nivel_riesgo", {})
    h = doc.add_paragraph()
    rh = h.add_run(f"{s.get('nombre', 'Sector')}  ·  NRI {nr.get('nivel')} ({nr.get('categoria')})")
    rh.bold = True
    rh.font.size = Pt(11)
    rh.font.color.rgb = ORANGE

    calc = s.get("calculo", {})
    actividad = s.get("actividad") or "—"
    modo_pred = "Almacenamiento" if s.get("modo_predominante") == "almacenamiento" else "Producción"
    pa = doc.add_paragraph()
    pa.add_run(f"Actividad: {actividad}  ·  Modo predominante: {modo_pred}").font.size = Pt(9)
    rows = []
    for z in calc.get("zonas", []):
        m = z.get("modo")
        modo = {"almacenamiento": "Almacenamiento", "deposito": "Depósito"}.get(m, "Producción")
        if m == "almacenamiento":
            dens, hv = f"{_fmt(z.get('qvi'))} MJ/m³", f"{_fmt(z.get('hi'))} m"
        elif m == "deposito":
            dens, hv = f"PCI {_fmt(z.get('pci'))}·ρ {_fmt(z.get('dens_comb'))}", f"{_fmt(z.get('volumen'))} m³"
        else:
            dens, hv = f"{_fmt(z.get('qsi'))} MJ/m²", "—"
        rows.append([z.get("descripcion", ""), modo, _fmt(z.get("ci"), 2), dens, hv,
                     f"{_fmt(z.get('superficie'))} m²", f"{_fmt(z.get('aporte'))} MJ"])
    totals = [f"Σ(qi·Ci·Si) = {_fmt(calc.get('numerador'))} MJ · A = {_fmt(calc.get('A'))} m² · Ra = {_fmt(calc.get('ra'), 2)}",
              "", "", "", "", "Qs =", f"{_fmt(calc.get('qs_mj'), 1)} MJ/m²"]
    _generic_table(doc, ["Actividad / zona", "Modo", "Ci", "qsi / qvi", "hi / Vol", "Superficie", "Parcial"], rows, totals=totals)

    sec = s.get("sectorizacion", {})
    if sec.get("mensaje"):
        doc.add_paragraph(sec["mensaje"])

    req = s.get("requisitos", {})
    _kv_table(doc, [
        {"k": "Estabilidad estructura", "v": (req.get("estabilidad") or {}).get("valor", "-")},
        {"k": "Delimitadores", "v": (req.get("delimitadores") or {}).get("valor", "-")},
        {"k": "Evacuación (máx.)", "v": f"{(s.get('evacuacion') or {}).get('distancia_maxima', '-')} m · "
                                        f"{(s.get('evacuacion') or {}).get('num_salidas', '-')} salidas"},
    ])

    for it in req.get("instalaciones", []):
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(f"{it.get('nombre')} — {'Obligatorio' if it.get('obligatorio') else 'No exigible'}")
        r.font.size = Pt(9)
        aero = it.get("superficie_aerodinamica")
        if aero:
            sub = doc.add_paragraph(style="List Bullet 2")
            sr = sub.add_run(f"Superficie aerodinámica de evacuación de humos (Anexo III, 8.4): "
                             f"≥ {_fmt(aero.get('superficie_aerodinamica'), 2)} m² "
                             f"(0,5 m² / {_fmt(aero.get('ratio_m2'), 0)} m², "
                             f"{'bajo rasante' if aero.get('situacion') == 'bajo_rasante' else 'sobre rasante'}).")
            sr.font.size = Pt(9)
            sr.italic = True

    for l in s.get("locales_riesgo_especial", []):
        p = doc.add_paragraph(style="List Bullet")
        magnitud = ""
        if l.get("volumen") is not None:
            magnitud = f"{_fmt(l.get('volumen'), 1)} m³ · "
        elif l.get("potencia") is not None:
            magnitud = f"{_fmt(l.get('potencia'), 1)} · "
        uso = f" · uso: {l['uso']}" if l.get("uso") else ""
        crit = f" ({l['criterio']})" if l.get("criterio") else ""
        r = p.add_run(f"LRE «{l.get('nombre')}» · riesgo {str(l.get('riesgo', '')).lower()} · "
                      f"{magnitud}{_fmt(l.get('superficie'))} m² · {l.get('resistencia')}{uso}{crit}"
                      f"{'. ' + l['justificacion'] if l.get('justificacion') else ''}")
        r.font.size = Pt(9)

    for e in s.get("escaleras_protegidas", []):
        p = doc.add_paragraph(style="List Bullet")
        tipo = ESC_TIPO_LABEL.get(e.get("tipo"), "protegida")
        sentido = f" · {e.get('sentido')}" if e.get("sentido") else ""
        c = e.get("calculo") or {}
        calc_txt = ""
        if c:
            calc_txt = (f" · {_fmt(c.get('ocupantes'))} ocupantes · altura {_fmt(c.get('altura_evacuacion'), 2)} m"
                        f" · proyectada {c.get('proteccion_proyectada_label', tipo)} (recinto {c.get('resistencia', 'EI 120')})"
                        f" · exige mín. {c.get('proteccion_requerida_label', 'no protegida')} (recinto {c.get('resistencia_requerida', '—')})"
                        f" · anchura mín. {c.get('anchura_min_texto', '—')}"
                        f" / proyectada {_fmt(c.get('anchura_proyectada'), 2)} m"
                        f" ({'CUMPLE' if c.get('estado') == 'cumple' else 'NO CUMPLE' if c.get('estado') == 'no_cumple' else 'pendiente'})")
        r = p.add_run(f"Escalera «{e.get('nombre')}» · {tipo}{sentido} · {e.get('situacion')}{calc_txt}"
                      f"{'. ' + e['justificacion'] if e.get('justificacion') else ''}")
        r.font.size = Pt(9)

    c = s.get("construccion")
    if c:
        p = doc.add_paragraph()
        pr = p.add_run(f"Construcción y propagación — {c.get('normativa', 'RSCIEI Anexo II')}")
        pr.bold = True
        pr.font.size = Pt(9.5)
        ip = doc.add_paragraph()
        ipr = ip.add_run(f"Propagación interior · {(c.get('interior') or {}).get('referencia', '')}")
        ipr.italic = True
        ipr.font.size = Pt(9)
        tbl = doc.add_table(rows=1, cols=3)
        tbl.style = "Table Grid"
        for i, h in enumerate(["Situación del elemento", "Techos y paredes", "Suelos"]):
            _set_cell(tbl.rows[0].cells[i], h, bold=True)
            _shade(tbl.rows[0].cells[i], "E2E8F0")
        for row in c["reaccion"]["tabla"]:
            cells = tbl.add_row().cells
            _set_cell(cells[0], row["situacion"])
            _set_cell(cells[1], row["techos_paredes"])
            _set_cell(cells[2], row["suelos"])
        rp = doc.add_paragraph(style="List Bullet")
        rp.add_run("Reacción al fuego de revestimientos: "
                   + ("el proyecto cumple las clases exigidas." if c["reaccion"]["cumple"] else "pendiente de verificar/justificar.")).font.size = Pt(9)
        delim = (c.get("interior") or {}).get("delimitadores")
        if delim:
            dp = doc.add_paragraph(style="List Bullet")
            dp.add_run(f"Elementos delimitadores del sector: exigido {delim.get('requerido', '—')} ({delim.get('detalle', '')}).").font.size = Pt(9)
        ep = doc.add_paragraph()
        epr = ep.add_run(f"Propagación exterior · {(c.get('exterior') or {}).get('referencia', '')}")
        epr.italic = True
        epr.font.size = Pt(9)
        if c.get("medianeras"):
            m = c["medianeras"]
            mp = doc.add_paragraph(style="List Bullet")
            mp.add_run(f"Medianería/separación con otro establecimiento/edificio: exigido {m['requerido']}, "
                       f"proyectado {m['proyectado'] or '—'} ({m['estado']}).").font.size = Pt(9)
        if c.get("fachadas"):
            f = c["fachadas"]
            fp = doc.add_paragraph(style="List Bullet")
            d_txt = f"d ≥ {f['distancia_d']} m, " if f.get("distancia_d") is not None else ""
            fp.add_run(f"Fachada en encuentro de sector: α {f['angulo']}°, {d_txt}"
                       f"exigido {f['requerido']}, proyectado {f['proyectado'] or '—'} ({f['estado']}).").font.size = Pt(9)
            h = f.get("huecos")
            if h:
                hp = doc.add_paragraph(style="List Bullet 2")
                proy = f"proyectada {h['proyectado_m']:.2f} m, " if h.get("proyectado_m") is not None else ""
                hp.add_run(f"Separación entre huecos (SI 2, Fig. 1.1): ángulo {h['angulo']:.0f}°, "
                           f"mínima {h['efectivo_m']:.2f} m, {proy}({h['estado']}).").font.size = Pt(9)
        cub = c["cubierta"]
        cp = doc.add_paragraph(style="List Bullet")
        cub_req = f" Exigencia: {cub['requerido']}." if cub.get("requerido") else ""
        cp.add_run("Cubierta " + ("ligera." if cub["ligera"] else "no ligera.") + cub_req).font.size = Pt(9)


    dbsi = s.get("dbsi")
    if dbsi:
        p = doc.add_paragraph()
        pr = p.add_run("CTE DB-SI · verificación por secciones")
        pr.bold = True
        pr.font.size = Pt(9.5)
        for sc in dbsi.get("secciones", []):
            est = sc.get("estado")
            label = "No cumple" if est == "no_cumple" else "Cumple" if est == "cumple" else "Info"
            b = doc.add_paragraph(style="List Bullet")
            rb = b.add_run(f"{sc.get('titulo', sc.get('id'))} — {label}")
            rb.font.size = Pt(9)
        for t in dbsi.get("tablas", []):
            _ref_table(doc, t)


def build_docx(report: dict, study: dict, project: dict) -> BytesIO:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)

    title = doc.add_paragraph()
    tr = title.add_run(report.get("title", "Informe"))
    tr.bold = True
    tr.font.size = Pt(18)
    tr.font.color.rgb = ORANGE
    sub = doc.add_paragraph()
    sr = sub.add_run(f"{study.get('name', '')}  ·  {project.get('code', '')} — {project.get('name', '')}")
    sr.font.color.rgb = GREY
    doc.add_paragraph(f"Generado: {report.get('generated_at', '')[:19].replace('T', ' ')}")

    for s in report.get("sections", []):
        stype = s.get("type")
        if stype == "ambito_header":
            p = doc.add_paragraph()
            pr = p.add_run(f"{s.get('tipo', '')}")
            pr.font.size = Pt(8)
            pr.font.color.rgb = GREY
            doc.add_heading(s.get("title", ""), level=1)
            continue

        doc.add_heading(s.get("title", ""), level=2)

        if stype == "text":
            doc.add_paragraph(s.get("body", ""))
        elif stype == "keyvalue":
            _kv_table(doc, s.get("rows", []))
        elif stype == "resumen_ejecutivo":
            rows = [[f.get("nombre"), f.get("ambito"), _fmt(f.get("qs_mj"), 1),
                     f"{f.get('nivel')} · {f.get('categoria')}",
                     {"cumple": "Cumple", "no_cumple": "No cumple"}.get(f.get("estado"), "Observaciones")]
                    for f in s.get("filas", [])]
            _generic_table(doc, ["Sector", "Ámbito", "Qs", "NRI", "Estado"], rows)
            for dep in s.get("depositos", []):
                p = doc.add_paragraph(style="List Bullet")
                r = p.add_run(dep.get("justificacion", ""))
                r.font.size = Pt(9)
        elif stype == "superficies":
            d = s.get("data", {})
            rows = [[f.get("nombre"), f.get("planta"), f.get("sector"),
                     _fmt(f.get("sup_util"), 2), _fmt(f.get("sup_construida"), 2)] for f in d.get("filas", [])]
            _generic_table(doc, ["Dependencia", "Planta", "Sector", "S. útil", "S. construida"], rows,
                           totals=["Total", "", "", _fmt(d.get("total_util"), 2), _fmt(d.get("total_construida"), 2)])
            grupos = d.get("grupos") or []
            if grupos:
                gp = doc.add_paragraph()
                gpr = gp.add_run("Superficies por planta")
                gpr.bold = True
                gpr.font.size = Pt(9.5)
                grows = [[g.get("planta"), _fmt(g.get("sub_util"), 2), _fmt(g.get("sub_construida"), 2)] for g in grupos]
                _generic_table(doc, ["Planta", "S. útil", "S. construida"], grows,
                               totals=["Total", _fmt(d.get("total_util"), 2), _fmt(d.get("total_construida"), 2)])
        elif stype == "ocupacion":
            d = s.get("data", {})
            rows = [[f.get("nombre"), f.get("sector"), f.get("uso"), _fmt(f.get("superficie"), 2),
                     str(f.get("densidad")), str(f.get("ocupacion"))] for f in d.get("filas", [])]
            _generic_table(doc, ["Dependencia", "Sector", "Uso", "Sup.", "Densidad", "Ocupación"], rows,
                           totals=["Ocupación total", "", "", "", "", str(d.get("total"))])
            for f in d.get("filas", []):
                if f.get("nota"):
                    p = doc.add_paragraph(style="List Bullet")
                    pr = p.add_run(f"Ocupación personalizada · {f.get('nombre')}: {f.get('nota')}")
                    pr.italic = True
                    pr.font.size = Pt(8)
            pp = d.get("por_planta") or []
            if pp:
                op = doc.add_paragraph()
                opr = op.add_run("Ocupación por planta")
                opr.bold = True
                opr.font.size = Pt(9.5)
                _generic_table(doc, ["Planta", "Ocupación"], [[x.get("planta"), str(x.get("ocupacion"))] for x in pp])
        elif stype == "resumen":
            rows = [[f.get("nombre"), f.get("config"), _fmt(f.get("superficie")), _fmt(f.get("qs_mj"), 1),
                     f"{f.get('nivel')} · {f.get('categoria')}", f.get("estabilidad"), str(f.get("ocupacion"))]
                    for f in s.get("filas", [])]
            _generic_table(doc, ["Sector", "Config.", "Sup.", "Qs", "NRI", "Estab.", "Ocup."], rows)
        elif stype == "sectores":
            for sec in s.get("items", []):
                _sector_docx(doc, sec)
        elif stype == "tablas":
            for t in s.get("items", []):
                _ref_table(doc, t)
        elif stype == "instalaciones":
            for it in s.get("items", []):
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(f"{it.get('nombre')} — {'Obligatorio' if it.get('obligatorio') else 'No exigible'}: {it.get('detalle', '')}")
        elif stype == "metricas":
            rows = [[m.get("label"), f"{m.get('valor')} {m.get('unidad', '')}".strip()] for m in s.get("items", [])]
            _generic_table(doc, ["Magnitud", "Valor"], rows)
        elif stype == "estancias_resumen":
            d = s.get("data", {})
            rows = [[f.get("nombre"), f.get("planta"), _fmt(f.get("sup_util"), 2), _fmt(f.get("sup_construida"), 2),
                     _fmt(f.get("volumen"), 2) if f.get("volumen") else "—", f.get("uso"), str(f.get("ocupacion"))]
                    for f in d.get("filas", [])]
            _generic_table(doc, ["Estancia", "Planta", "S. útil", "S. constr.", "Vol.", "Uso", "Ocup."], rows,
                           totals=["Total", "", _fmt(d.get("total_util"), 2), _fmt(d.get("total_construida"), 2),
                                   _fmt(d.get("total_volumen"), 2), "", str(d.get("total_ocupacion"))])
        elif stype == "secciones_modulo":
            for sec in s.get("items", []):
                est = sec.get("estado")
                label = {"cumple": "Cumple", "no_cumple": "No cumple", "observaciones": "Revisar"}.get(est, "Info")
                p = doc.add_paragraph()
                pr = p.add_run(f"{sec.get('titulo', '')}  —  {label}")
                pr.bold = True
                pr.font.size = Pt(10)
                if sec.get("descripcion"):
                    dp = doc.add_paragraph()
                    dr = dp.add_run(sec["descripcion"])
                    dr.italic = True
                    dr.font.size = Pt(8.5)
                    dr.font.color.rgb = GREY
                for it in sec.get("items", []):
                    b = doc.add_paragraph(style="List Bullet")
                    b.add_run(f"{it.get('nombre')}: {it.get('detalle', '')}").font.size = Pt(9)
        elif stype == "rebt_potencias":
            rows = [[r.get("cuadro"), r.get("circuito"), r.get("tipo"), _fmt(r.get("potencia"), 0),
                     str(r.get("factor_uso")), _fmt(r.get("potencia_calculo"), 0)] for r in s.get("items", [])]
            _generic_table(doc, ["Cuadro", "Circuito", "Tipo", "P inst. (W)", "F.u.", "P cálc. (W)"], rows,
                           totals=["TOTAL", "", "", _fmt(s.get("total_instalada"), 0), "", _fmt(s.get("total_calculo"), 0)])
        elif stype == "rebt_lineas":
            lbl = {"cumple": "Cumple", "no_cumple": "No cumple", "observaciones": "Revisar"}
            rows = [[l.get("cuadro"), l.get("circuito"), _fmt(l.get("Ib"), 2), l.get("seccion_txt"),
                     _fmt(l.get("iz"), 0), l.get("proteccion"),
                     ("Sí" if l.get("protege_cable") else "Revisar") if l.get("coordinacion") else "—",
                     f"{l.get('caida_acum_pct')}%", l.get("diferencial"), l.get("tierra"),
                     lbl.get(l.get("estado"), "Info")] for l in s.get("items", [])]
            _generic_table(doc, ["Cuadro", "Línea", "Ib", "Sección", "Iz", "Protec.", "Protege cable", "Caída ac.", "Difer.", "Tierra", "Estado"], rows)
            nota = doc.add_paragraph().add_run("Coordinación protección↔cable (ITC-BT-22): Ib ≤ In ≤ Iz e I2 ≤ 1,45·Iz.")
            nota.italic = True
            nota.font.size = Pt(8)
            nota.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        elif stype == "rebt_cortocircuito":
            yn = lambda ok: "Sí" if ok else "No"
            rows = [[l.get("cuadro"), l.get("circuito"),
                     _fmt((l.get("cortocircuito") or {}).get("icc_max_a"), 0),
                     _fmt((l.get("cortocircuito") or {}).get("icc_min_a"), 0),
                     _fmt((l.get("cortocircuito") or {}).get("im_a"), 0),
                     yn((l.get("cortocircuito") or {}).get("pdc_ok")),
                     yn((l.get("cortocircuito") or {}).get("termico_ok")),
                     yn((l.get("cortocircuito") or {}).get("magnetico_ok")),
                     yn((l.get("cortocircuito") or {}).get("protege_cc")),
                     ((l.get("cortocircuito") or {}).get("recomendacion") or "—").replace("Recomendación: ", "")] for l in s.get("items", [])]
            _generic_table(doc, ["Cuadro", "Línea", "Icc máx (A)", "Icc mín (A)", "Im (A)", "P. corte", "Térmico k²S²≥I²t", "Disparo Icc≥Im", "Protege CC", "Recomendación"], rows)
            nota2 = doc.add_paragraph().add_run("Poder de corte ≥ Icc máx · k²·S² ≥ I²t · Icc mín ≥ Im (disparo magnético). Icc por método simplificado (Guía BT-22), orientativo.")
            nota2.italic = True
            nota2.font.size = Pt(8)
            nota2.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        elif stype == "rebt_anexo_calculo":
            estado_lbl = {"cumple": "Cumple", "no_cumple": "No cumple", "observaciones": "Revisar"}
            for b in s.get("items", []):
                hp = doc.add_paragraph()
                hr = hp.add_run(f"{b.get('circuito', 'Circuito')} — {estado_lbl.get(b.get('estado'), 'Info')}")
                hr.bold = True
                hr.font.size = Pt(9.5)
                for p in b.get("pasos", []):
                    tp = doc.add_paragraph(style="List Bullet")
                    tp.add_run(f"{p.get('tema')}: ").bold = True
                    tp.runs[0].font.size = Pt(9)
                    r2 = tp.add_run(f"{p.get('formula')} → {p.get('valor')} ")
                    r2.font.size = Pt(9)
                    r3 = tp.add_run(f"({p.get('referencia')})")
                    r3.italic = True
                    r3.font.size = Pt(8)
                    r3.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
                doc.add_paragraph()
        elif stype == "rebt_esquema":
            def _board_docx(b, nivel):
                spd = b.get("spd") or {}
                head = f"■ {b.get('nombre')} · IGA {b.get('iga')} A · ID {b.get('diferencial_general_ma')} mA"
                if spd.get("presente"):
                    head += f" · SPD {spd.get('tipo')}"
                head += f" · P cálc. {b.get('potencia_demandada')} W · {(b.get('alimentacion') or {}).get('seccion_txt')} mm²"
                p = doc.add_paragraph()
                pr = p.add_run(("    " * nivel) + head)
                pr.bold = True
                pr.font.size = Pt(9)
                for n in b.get("nodos", []):
                    k = n.get("kind")
                    if k == "linea":
                        cp = doc.add_paragraph(style="List Bullet")
                        cp.add_run(("    " * nivel) + f"{n.get('nombre')} · {n.get('magneto')} A · {n.get('sensibilidad')} mA · "
                                   f"{n.get('seccion_txt')} mm² · {n.get('potencia')} W · caída {n.get('caida_acum_pct')}% "
                                   f"(máx {n.get('limite_total')}%) · "
                                   f"{'CUMPLE' if n.get('estado') == 'cumple' else 'NO CUMPLE' if n.get('estado') == 'no_cumple' else 'revisar'}").font.size = Pt(8.5)
                    elif k == "puente":
                        pp = doc.add_paragraph()
                        ppr = pp.add_run(("    " * nivel) + f"↳ {n.get('nombre')} (línea puente) · {n.get('magneto')} A · "
                                         f"{n.get('sensibilidad')} mA · {n.get('seccion_txt')} mm² · caída {n.get('caida_acum_pct')}%")
                        ppr.italic = True
                        ppr.font.size = Pt(8.5)
                        for sub in n.get("sublineas", []):
                            sp = doc.add_paragraph(style="List Bullet")
                            sp.add_run(("    " * (nivel + 1)) + f"{sub.get('nombre')} · {sub.get('magneto')} A · "
                                       f"{sub.get('seccion_txt')} mm² · {sub.get('potencia')} W · caída {sub.get('caida_acum_pct')}% "
                                       f"(máx {sub.get('limite_total')}%)").font.size = Pt(8)
                    elif k == "subcuadro":
                        _board_docx(n.get("board") or {}, nivel + 1)
            if s.get("arbol"):
                _board_docx(s["arbol"], 0)
        elif stype == "observaciones":
            for o in s.get("items", []):
                doc.add_paragraph(o.get("texto", ""), style="List Bullet")
        elif stype == "anexo_propagacion":
            for b in s.get("items", []):
                hp = doc.add_paragraph()
                hr = hp.add_run(f"{b.get('sector', 'Sector')} — {b.get('normativa', '')}")
                hr.bold = True
                hr.font.size = Pt(10)
                for titulo, key in (("Propagación interior", "interior"), ("Propagación exterior", "exterior")):
                    sp = doc.add_paragraph()
                    spr = sp.add_run(titulo)
                    spr.italic = True
                    spr.bold = True
                    spr.font.size = Pt(9)
                    for f in b.get(key, []):
                        tp = doc.add_paragraph()
                        tp.add_run(f.get("tema", "")).bold = True
                        tp.runs[0].font.size = Pt(9)
                        cp = doc.add_paragraph()
                        cr = cp.add_run(f.get("cita", ""))
                        cr.italic = True
                        cr.font.size = Pt(8.5)
                        rp = doc.add_paragraph()
                        rr = rp.add_run(f"— {f.get('referencia', '')}")
                        rr.font.size = Pt(8)
                        rr.font.color.rgb = GREY
                        if f.get("valor"):
                            vp = doc.add_paragraph(style="List Bullet")
                            vp.add_run(f.get("valor")).font.size = Pt(8.5)
                doc.add_paragraph()
        elif stype == "fundamentos":
            for fnd in s.get("items", []):
                p = doc.add_paragraph()
                pr = p.add_run(fnd.get("tema", ""))
                pr.bold = True
                pr.font.size = Pt(9.5)
                cp = doc.add_paragraph()
                cr = cp.add_run(fnd.get("cita", ""))
                cr.italic = True
                cr.font.size = Pt(9)
                rp = doc.add_paragraph()
                rr = rp.add_run(f"— {fnd.get('referencia', '')}")
                rr.font.size = Pt(8)
                rr.font.color.rgb = GREY

    doc.add_paragraph()
    disc = doc.add_paragraph()
    dr = disc.add_run(report.get("disclaimer", ""))
    dr.italic = True
    dr.font.size = Pt(8)
    dr.font.color.rgb = GREY

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# --------------------------------------------------------------------------- #
# XLSX (con fórmulas vivas)
# --------------------------------------------------------------------------- #
HEAD_FILL = PatternFill("solid", fgColor="1E293B")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=10)
TOTAL_FILL = PatternFill("solid", fgColor="F1F5F9")
CALC_FILL = PatternFill("solid", fgColor="FEF3C7")
THIN = Side(style="thin", color="CBD5E1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _sheet_title(name):
    # Excel: máx 31 chars, sin caracteres inválidos
    for ch in "[]:*?/\\":
        name = name.replace(ch, " ")
    return name[:31]


def _header_row(ws, row, headers):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER


def _autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _build_generic_xlsx(wb: Workbook, results: dict, project: dict) -> BytesIO:
    from normative_modules.shared_estancias import resumen_estancias
    # Hoja Resumen
    ws = wb.active
    ws.title = "Resumen"
    ws.cell(row=1, column=1, value=f"Resumen · {project.get('code','')} {project.get('name','')}").font = Font(bold=True, size=12)
    _header_row(ws, 3, ["Magnitud", "Valor", "Unidad"])
    r = 4
    for m in results.get("resumen", []):
        ws.cell(row=r, column=1, value=m.get("label"))
        ws.cell(row=r, column=2, value=m.get("valor"))
        ws.cell(row=r, column=3, value=m.get("unidad"))
        for c in range(1, 4):
            ws.cell(row=r, column=c).border = BORDER
        r += 1
    ws.cell(row=r + 1, column=1, value="Estado global").font = Font(bold=True)
    ws.cell(row=r + 1, column=2, value=results.get("estado_global")).font = Font(bold=True)
    _autosize(ws, [34, 20, 12])

    # Hoja Estancias
    est = resumen_estancias(results.get("estancias") or [])
    if est["filas"]:
        we = wb.create_sheet("Estancias")
        we.cell(row=1, column=1, value="Cuadro de superficies y ocupación").font = Font(bold=True, size=12)
        _header_row(we, 3, ["Estancia", "Planta", "S. útil (m²)", "S. constr. (m²)", "Volumen (m³)", "Uso", "Ocupación"])
        r = 4
        for f in est["filas"]:
            for c, key in enumerate(["nombre", "planta", "sup_util", "sup_construida", "volumen", "uso", "ocupacion"], start=1):
                we.cell(row=r, column=c, value=f.get(key))
                we.cell(row=r, column=c).border = BORDER
            r += 1
        we.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
        we.cell(row=r, column=3, value=est["total_util"]).font = Font(bold=True)
        we.cell(row=r, column=4, value=est["total_construida"]).font = Font(bold=True)
        we.cell(row=r, column=5, value=est["total_volumen"]).font = Font(bold=True)
        we.cell(row=r, column=7, value=est["total_ocupacion"]).font = Font(bold=True)
        for c in range(1, 8):
            we.cell(row=r, column=c).fill = TOTAL_FILL
            we.cell(row=r, column=c).border = BORDER
        _autosize(we, [28, 14, 16, 16, 14, 22, 14])

    # Hoja Verificación (RITE) o Cuadro de potencias + Líneas (REBT)
    if results.get("lineas") is not None:
        wp = wb.create_sheet("Cuadro de potencias")
        _header_row(wp, 1, ["Cuadro", "Circuito", "Tipo", "P instalada (W)", "F. uso", "P cálculo (W)"])
        r = 2
        for it in results.get("cuadro_potencias", []):
            for c, key in enumerate(["cuadro", "circuito", "tipo", "potencia", "factor_uso", "potencia_calculo"], start=1):
                wp.cell(row=r, column=c, value=it.get(key))
                wp.cell(row=r, column=c).border = BORDER
            r += 1
        wp.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
        wp.cell(row=r, column=4, value=results.get("datos", {}).get("potencia_instalada")).font = Font(bold=True)
        wp.cell(row=r, column=6, value=results.get("datos", {}).get("potencia_demandada")).font = Font(bold=True)
        _autosize(wp, [26, 26, 14, 16, 10, 16])

        wl = wb.create_sheet("Líneas")
        _header_row(wl, 1, ["Cuadro", "Línea", "Ib (A)", "Sección (mm²)", "Iz (A)", "Caída acum. (%)", "Protección", "Diferencial", "Tierra (mm²)", "Estado"])
        r = 2
        lbl = {"cumple": "Cumple", "no_cumple": "No cumple", "observaciones": "Revisar"}
        for l in results.get("lineas", []):
            vals = [l.get("cuadro"), l.get("circuito"), l.get("Ib"), l.get("seccion_txt"), l.get("iz"),
                    l.get("caida_acum_pct"), l.get("proteccion"), l.get("diferencial"), l.get("tierra"), lbl.get(l.get("estado"), "Info")]
            for c, v in enumerate(vals, start=1):
                wl.cell(row=r, column=c, value=v)
                wl.cell(row=r, column=c).border = BORDER
            r += 1
        _autosize(wl, [24, 30, 10, 14, 10, 16, 14, 16, 12, 12])
    else:
        wv = wb.create_sheet("Verificación")
        wv.cell(row=1, column=1, value="Verificación de exigencias").font = Font(bold=True, size=12)
        _header_row(wv, 3, ["Apartado", "Estado", "Concepto", "Detalle"])
        r = 4
        labels = {"cumple": "Cumple", "no_cumple": "No cumple", "observaciones": "Revisar", "info": "Info"}
        for sec in results.get("secciones", []):
            for it in sec.get("items", []):
                wv.cell(row=r, column=1, value=sec.get("titulo"))
                wv.cell(row=r, column=2, value=labels.get(sec.get("estado"), "Info"))
                wv.cell(row=r, column=3, value=it.get("nombre"))
                wv.cell(row=r, column=4, value=it.get("detalle"))
                for c in range(1, 5):
                    wv.cell(row=r, column=c).border = BORDER
                r += 1
        _autosize(wv, [34, 12, 32, 60])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_xlsx(study: dict, project: dict) -> BytesIO:
    results = study.get("results") or {}
    wb = Workbook()

    # Módulos genéricos (REBT/RITE): sin sectores de incendio → libro genérico.
    if (results.get("secciones") is not None or results.get("lineas") is not None) and not results.get("sectores"):
        return _build_generic_xlsx(wb, results, project)

    # --- Hoja Superficies --------------------------------------------------- #
    ws = wb.active
    ws.title = "Superficies"
    ws.cell(row=1, column=1, value=f"Cuadro de superficies · {project.get('code','')} {project.get('name','')}").font = Font(bold=True, size=12)
    _header_row(ws, 3, ["Dependencia", "Planta", "Sector", "S. útil (m²)", "S. construida (m²)"])
    sup = results.get("superficies", {})
    r = 4
    first = r
    for f in sup.get("filas", []):
        ws.cell(row=r, column=1, value=f.get("nombre"))
        ws.cell(row=r, column=2, value=f.get("planta"))
        ws.cell(row=r, column=3, value=f.get("sector"))
        ws.cell(row=r, column=4, value=f.get("sup_util"))
        ws.cell(row=r, column=5, value=f.get("sup_construida"))
        for c in range(1, 6):
            ws.cell(row=r, column=c).border = BORDER
        r += 1
    last = r - 1
    tot = ws.cell(row=r, column=1, value="TOTAL")
    tot.font = Font(bold=True)
    if last >= first:
        ws.cell(row=r, column=4, value=f"=SUM(D{first}:D{last})").font = Font(bold=True)
        ws.cell(row=r, column=5, value=f"=SUM(E{first}:E{last})").font = Font(bold=True)
    for c in range(1, 6):
        ws.cell(row=r, column=c).fill = TOTAL_FILL
        ws.cell(row=r, column=c).border = BORDER
    grupos = sup.get("grupos") or []
    if grupos:
        r += 2
        ws.cell(row=r, column=1, value="Superficies por planta").font = Font(bold=True)
        r += 1
        _header_row(ws, r, ["Planta", "", "", "S. útil (m²)", "S. construida (m²)"])
        r += 1
        for g in grupos:
            ws.cell(row=r, column=1, value=g.get("planta"))
            ws.cell(row=r, column=4, value=g.get("sub_util"))
            ws.cell(row=r, column=5, value=g.get("sub_construida"))
            for c in range(1, 6):
                ws.cell(row=r, column=c).border = BORDER
            r += 1
    _autosize(ws, [30, 14, 22, 16, 18])

    # --- Hoja Ocupación ----------------------------------------------------- #
    wo = wb.create_sheet("Ocupación")
    wo.cell(row=1, column=1, value="Cálculo de ocupación (CTE DB-SI, Tabla 2.1)").font = Font(bold=True, size=12)
    _header_row(wo, 3, ["Dependencia", "Sector", "Uso", "Superficie (m²)", "Densidad", "Ocupación", "Nota"])
    occ = results.get("ocupacion", {})
    r = 4
    for f in occ.get("filas", []):
        wo.cell(row=r, column=1, value=f.get("nombre"))
        wo.cell(row=r, column=2, value=f.get("sector"))
        wo.cell(row=r, column=3, value=f.get("uso"))
        wo.cell(row=r, column=4, value=f.get("superficie"))
        wo.cell(row=r, column=5, value=str(f.get("densidad")))
        wo.cell(row=r, column=6, value=f.get("ocupacion"))
        wo.cell(row=r, column=7, value=f.get("nota") or "")
        for c in range(1, 8):
            wo.cell(row=r, column=c).border = BORDER
        r += 1
    wo.cell(row=r, column=1, value="OCUPACIÓN TOTAL").font = Font(bold=True)
    wo.cell(row=r, column=6, value=occ.get("total")).font = Font(bold=True)
    for c in range(1, 8):
        wo.cell(row=r, column=c).fill = TOTAL_FILL
        wo.cell(row=r, column=c).border = BORDER
    porp = occ.get("por_planta") or []
    if porp:
        r += 2
        wo.cell(row=r, column=1, value="Ocupación por planta").font = Font(bold=True)
        r += 1
        for x in porp:
            wo.cell(row=r, column=1, value=x.get("planta"))
            wo.cell(row=r, column=6, value=x.get("ocupacion"))
            for c in range(1, 8):
                wo.cell(row=r, column=c).border = BORDER
            r += 1
    _autosize(wo, [30, 22, 22, 16, 16, 14, 40])

    # --- Hoja por sector: carga de fuego Qs con fórmulas -------------------- #
    for s in results.get("sectores", []):
        calc = s.get("calculo", {})
        zonas = calc.get("zonas", [])
        ws = wb.create_sheet(_sheet_title(f"Qs · {s.get('nombre', 'Sector')}"))
        nr = s.get("nivel_riesgo", {})
        ws.cell(row=1, column=1, value=f"Carga de fuego · {s.get('nombre','Sector')}").font = Font(bold=True, size=12)
        ws.cell(row=2, column=1, value="Los aportes y la Qs son fórmulas vivas: edite qsi/qvi/Ci/hi/superficie y Excel recalcula.").font = Font(italic=True, size=9, color="64748B")
        headers = ["Zona", "Modo", "qsi (MJ/m²)", "qvi (MJ/m³)", "Ci", "hi (m)", "Superficie (m²)", "Aporte (MJ)"]
        _header_row(ws, 4, headers)
        first = 5
        r = first
        for z in zonas:
            modo = z.get("modo")
            almac = modo == "almacenamiento"
            deposito = modo == "deposito"
            ws.cell(row=r, column=1, value=z.get("descripcion"))
            ws.cell(row=r, column=2, value={"almacenamiento": "Almacenamiento", "deposito": "Depósito"}.get(modo, "Producción"))
            ws.cell(row=r, column=3, value=(None if (almac or deposito) else z.get("qsi")))
            ws.cell(row=r, column=4, value=(z.get("qvi") if almac else None))
            ws.cell(row=r, column=5, value=z.get("ci"))
            ws.cell(row=r, column=6, value=(z.get("hi") if almac else None))
            ws.cell(row=r, column=7, value=z.get("superficie"))
            # Aporte: producción = qsi*Ci*S ; almacenamiento = qvi*Ci*hi*S ; depósito = PCI*densidad*volumen*Ci (valor calculado)
            if deposito:
                ws.cell(row=r, column=8, value=z.get("aporte"))
            elif almac:
                ws.cell(row=r, column=8, value=f"=D{r}*E{r}*F{r}*G{r}")
            else:
                ws.cell(row=r, column=8, value=f"=C{r}*E{r}*G{r}")
            for c in range(1, 9):
                ws.cell(row=r, column=c).border = BORDER
            r += 1
        last = r - 1
        r += 1
        a_row, ra_row, num_row, qs_row = r, r + 1, r + 2, r + 3
        ws.cell(row=a_row, column=1, value="Superficie útil del sector A (m²)").font = Font(bold=True)
        ws.cell(row=a_row, column=8, value=(f"=SUM(G{first}:G{last})" if last >= first else calc.get("A")))
        ws.cell(row=ra_row, column=1, value="Coeficiente de activación Ra").font = Font(bold=True)
        ws.cell(row=ra_row, column=8, value=calc.get("ra"))
        ws.cell(row=num_row, column=1, value="Suma de aportes (MJ)").font = Font(bold=True)
        ws.cell(row=num_row, column=8, value=(f"=SUM(H{first}:H{last})" if last >= first else calc.get("numerador")))
        ws.cell(row=qs_row, column=1, value="Qs = (Σ aportes / A) · Ra  (MJ/m²)").font = Font(bold=True, size=11)
        qs = ws.cell(row=qs_row, column=8, value=f"=IF(H{a_row}=0,0,H{num_row}/H{a_row}*H{ra_row})")
        qs.font = Font(bold=True, size=11)
        for rr in (a_row, ra_row, num_row, qs_row):
            for c in (1, 8):
                ws.cell(row=rr, column=c).fill = CALC_FILL if rr == qs_row else TOTAL_FILL
                ws.cell(row=rr, column=c).border = BORDER
        ws.cell(row=qs_row + 2, column=1,
                value=f"NRI resultante: nivel {nr.get('nivel')} ({nr.get('categoria')}). "
                      "Las escaleras y los locales de riesgo especial están excluidos del cálculo de la carga de fuego.").font = Font(italic=True, size=9, color="64748B")
        _autosize(ws, [28, 16, 14, 14, 10, 10, 16, 16])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
