"""Exportación del cuadro maestro de superficies/ocupación del expediente (XLSX y PDF)."""
import math
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

_HEAD = ["Estancia", "S. útil (m²)", "S. constr. (m²)", "Altura (m)", "Uso", "m²/pers.", "Ocup.", "Planta"]


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _ocupacion(e: dict):
    if e.get("sin_ocupacion"):
        return 0
    if e.get("ocupacion_custom") and e.get("ocupacion_manual") not in ("", None):
        return math.ceil(_num(e.get("ocupacion_manual")))
    dens = _num(e.get("densidad"))
    if dens <= 0:
        return None
    sup = _num(e.get("sup_util")) or _num(e.get("sup_construida"))
    return 0 if sup <= 0 else math.ceil(sup / dens)


def _rows(project: dict):
    estancias = project.get("estancias", []) or []
    plantas = [p.get("nombre") for p in (project.get("plantas", []) or [])]
    ordered = sorted(estancias, key=lambda e: (plantas.index(e["planta"]) if e.get("planta") in plantas else len(plantas)))
    rows = []
    for e in ordered:
        oc = _ocupacion(e)
        rows.append([
            e.get("nombre") or "", _num(e.get("sup_util")), _num(e.get("sup_construida")),
            _num(e.get("altura")), e.get("uso_ocupacion") or "", _num(e.get("densidad")) or "",
            "—" if oc is None else oc, e.get("planta") or "Sin planta",
        ])
    return rows


def _totales(rows):
    tu = sum(r[1] for r in rows)
    tc = sum(r[2] for r in rows)
    to = sum(r[6] for r in rows if isinstance(r[6], (int, float)))
    return tu, tc, to


def master_table_xlsx(project: dict) -> BytesIO:
    rows = _rows(project)
    tu, tc, to = _totales(rows)
    wb = Workbook()
    ws = wb.active
    ws.title = "Cuadro de superficies"
    ws.append([f"Cuadro de superficies y ocupación — {project.get('name', '')} ({project.get('code', '')})"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(_HEAD))
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(_HEAD)
    hdr = ws[2]
    for c in hdr:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F2937")
        c.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(r)
    ws.append([f"TOTAL ({len(rows)} estancias)", round(tu, 2), round(tc, 2), "", "", "", to, ""])
    total_row = ws[ws.max_row]
    for c in total_row:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E5E7EB")
    widths = [30, 14, 16, 12, 22, 12, 10, 18]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def master_table_pdf(project: dict) -> BytesIO:
    rows = _rows(project)
    tu, tc, to = _totales(rows)
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=15 * mm, rightMargin=15 * mm,
                            topMargin=15 * mm, bottomMargin=15 * mm)
    styles = getSampleStyleSheet()
    title = Paragraph(f"<b>Cuadro de superficies y ocupación</b><br/>{project.get('name', '')} ({project.get('code', '')})", styles["Title"])
    data = [_HEAD]
    for r in rows:
        data.append([str(x) if not isinstance(x, float) else f"{x:g}" for x in r])
    data.append([f"TOTAL ({len(rows)} estancias)", f"{tu:g}", f"{tc:g}", "", "", "", str(to), ""])
    tbl = Table(data, repeatRows=1)
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2937")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9CA3AF")),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E5E7EB")),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("ALIGN", (1, 0), (3, -1), "RIGHT"),
        ("ALIGN", (5, 0), (6, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F3F4F6")]),
    ]))
    doc.build([title, Spacer(1, 6 * mm), tbl])
    buf.seek(0)
    return buf
